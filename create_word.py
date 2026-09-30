from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_word_report():
    doc = Document()
    
    # Title
    title = doc.add_heading('Proposal: Integrating Audio for Low-Light Anomaly Detection', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Intro
    doc.add_heading('1. Introduction', level=1)
    doc.add_paragraph(
        "While our Spatial-Temporal Graph Convolutional Network (STGCN) model performs well for visual anomaly detection, it faces practical limitations in low-light environments, such as hospital wards at night. To ensure continuous and reliable monitoring, we propose integrating an audio analysis pipeline. This will allow the system to detect anomalies—such as falls or distress calls—through sound when video visibility is poor."
    )
    
    # Section 1
    doc.add_heading('2. Audio Feature Extraction', level=1)
    doc.add_paragraph("We will analyze the audio extracted from the video feeds using two primary methods:")
    
    p1 = doc.add_paragraph(style='List Bullet')
    p1.add_run("Volume Thresholding: ").bold = True
    p1.add_run("By measuring the Root Mean Square (RMS) energy or decibel (dB) levels, we can quickly detect sudden, loud noises (e.g., a crash or a shout) that stand out from the quiet baseline of a night ward.")
    
    p2 = doc.add_paragraph(style='List Bullet')
    p2.add_run("Sound Classification: ").bold = True
    p2.add_run("To distinguish between normal loud noises (like a door closing) and actual anomalies (like a person falling), we will extract Mel-Frequency Cepstral Coefficients (MFCCs). These features act as acoustic signatures, helping us classify the specific type of sound.")

    # Section 2
    doc.add_heading('3. Proposed Technology Stack', level=1)
    doc.add_paragraph("We will use established open-source tools to implement this efficiently:")
    
    p3 = doc.add_paragraph(style='List Bullet')
    p3.add_run("FFmpeg: ").bold = True
    p3.add_run("A reliable tool to extract the audio tracks from our existing video datasets. This allows us to use our current data without requiring new hardware.")
    
    p4 = doc.add_paragraph(style='List Bullet')
    p4.add_run("Librosa: ").bold = True
    p4.add_run("A standard Python library for audio processing. We will use it to compute decibel levels and extract MFCC features from the audio files.")

    # Section 3
    doc.add_heading('4. Multimodal Integration Strategy', level=1)
    doc.add_paragraph(
        "To combine the new audio system with our existing STGCN model, we recommend a 'Late Fusion' approach. This ensures our current video model remains unchanged while acting as a safety net."
    )
    
    p5 = doc.add_paragraph(style='List Number')
    p5.add_run("Independent Processing: ").bold = True
    p5.add_run("The STGCN model will generate a visual anomaly score, while the audio pipeline will independently generate an audio anomaly score.")
    
    p6 = doc.add_paragraph(style='List Number')
    p6.add_run("Dynamic Weighting: ").bold = True
    p6.add_run("We will implement a decision-making layer that assesses the lighting conditions. If the video frame is too dark, the system will dynamically increase the weight of the audio score to make the final prediction.")

    # Section 4
    doc.add_heading('5. Implementation Steps', level=1)
    doc.add_paragraph("We plan to execute this integration through the following steps:")
    
    doc.add_paragraph("Data Extraction: Use FFmpeg to isolate audio files from the current video dataset.", style='List Number')
    doc.add_paragraph("Baseline Analysis: Use Librosa to analyze decibel levels of normal versus anomalous events to establish initial volume thresholds.", style='List Number')
    doc.add_paragraph("Model Training (Optional): If volume thresholding is insufficient, train a lightweight classifier (such as a Random Forest) on the MFCC features to improve sound classification accuracy.", style='List Number')
    doc.add_paragraph("System Fusion: Develop the logic to combine the STGCN visual scores and the audio scores into a final, reliable prediction.", style='List Number')

    doc.add_paragraph("\nThis approach offers a practical, robust solution to the low-light limitation without overhauling our current architecture.")
    
    doc.save('e:/Capstone_Project/audio_integration_report.docx')

if __name__ == "__main__":
    create_word_report()
